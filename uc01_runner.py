#!/usr/bin/env python3
"""
UC-01 Runner: Introspección con Fibonacci Extendido
Busca el punto de quiebre donde el servidor ya no puede responder.
Secuencia Fibonacci: 1, 2, 3, 5, 8, 13, 21, 34, 55, 89, 144, 233, 377
Genera Excel y gráficos en pruebas/UC01_Introspeccion/
"""
import os, sys, time, subprocess, re, glob
from datetime import datetime

# --- Dependencias opcionales ---
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

# ───────────────────────────────────────────────────────────────────────
# CONFIG
# ───────────────────────────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE_DIR, "pruebas", "UC01_Introspeccion")
os.makedirs(OUT_DIR, exist_ok=True)

ENVS = {
    "Vulnerable": os.path.join(BASE_DIR, "E-commerce"),
    "Protegido":  os.path.join(BASE_DIR, "E-commerce-controles"),
}

UC_SCRIPT = "UC01_introspeccion/uc01_stress_introspeccion.js"
RUNS = 3

# Fibonacci extendido: va subiendo hasta encontrar el punto de quiebre
FIBONACCI_VUS = [1, 2, 3, 5, 8, 13, 21, 34, 55, 89, 144, 233, 377]

# Umbral de fallos para declarar "punto de quiebre"
FAIL_THRESHOLD = 50.0  # Si los fallos superan 50%, el servidor colapsó

CONTAINERS = ['api-gateway', 'mongo-db', 'ms-catalogo', 'ms-ordenes', 
              'ms-resenas', 'ms-usuarios', 'postgres-db']

# ───────────────────────────────────────────────────────────────────────
# DOCKER HELPERS
# ───────────────────────────────────────────────────────────────────────
def restart_env_full(env_path):
    """Reinicio completo de entornos para liberar recursos."""
    print(f"\n[+] Full Restart de {env_path}...")
    subprocess.run(["docker-compose", "down", "--remove-orphans"], cwd=ENVS["Vulnerable"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
    subprocess.run(["docker-compose", "down", "--remove-orphans"], cwd=ENVS["Protegido"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
    for _retry in range(3):
        # Nuke known containers just in case
        subprocess.run(["docker", "rm", "-f", "exp-vulnerable-api-gateway-1", "exp-vulnerable-mongo-db-1", "exp-vulnerable-postgres-db-1", "exp-vulnerable-ms-usuarios-1", "exp-vulnerable-ms-ordenes-1", "exp-vulnerable-ms-catalogo-1", "exp-vulnerable-ms-resenas-1", "exp-vulnerable-mongo-init-1"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        time.sleep(5) # Let Docker Desktop catch up on Windows
        result = subprocess.run(["docker-compose", "up", "-d"], cwd=env_path, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=600)
        if result.returncode == 0:
            break
        print(f"[WARN] docker-compose up falló (intento {_retry+1}/3). Reintentando...")
        subprocess.run(["docker-compose", "down"], cwd=env_path, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
        time.sleep(3)
    else:
        print(f"[ERROR] docker-compose up falló tras 3 intentos: {result.stderr.decode('utf-8', errors='ignore')}")
        sys.exit(1)
    # Verificación activa: esperar hasta que el gateway responda
    import urllib.request, urllib.error
    print("[HEALTHCHECK] Verificando que el gateway responda antes de continuar...")
    for _attempt in range(30):  # max 60s (30 x 2s)
        time.sleep(2)
        try:
            req = urllib.request.Request("http://localhost:4000/graphql",
                data=b'{"query":"{ __typename }"}',
                headers={"Content-Type": "application/json"}, method="POST")
            resp = urllib.request.urlopen(req, timeout=5)
            print(f"[OK] Gateway respondió (status={resp.status}) tras {(_attempt+1)*2}s")
            break
        except urllib.error.HTTPError as e:
            if e.code in (400, 403):
                print(f"[OK] Gateway respondió (status={e.code}) tras {(_attempt+1)*2}s")
                break
        except Exception:
            pass
    else:
        print("[WARN] Gateway no respondió tras 60s. Continuando de todas formas...")
    # Para Vulnerable: verificar además que la federación esté lista (introspección funcional)
    if env_path == ENVS["Vulnerable"]:
        print("[HEALTHCHECK-FEDERATION] Verificando que la federación de subgrafos esté completa...")
        for _attempt in range(15):  # max 30s
            time.sleep(2)
            try:
                req = urllib.request.Request("http://localhost:4000/graphql",
                    data=b'{"query":"{ __schema { queryType { name } } }"}',
                    headers={"Content-Type": "application/json"}, method="POST")
                resp = urllib.request.urlopen(req, timeout=5)
                body = resp.read().decode()
                if "__schema" in body and "errors" not in body:
                    print(f"[OK] Federación lista: introspección funcional tras {(_attempt+1)*2}s extra")
                    break
            except Exception:
                pass
        else:
            print("[WARN] Federación no verificada tras 30s extra.")
    time.sleep(8)  # Margen extra de estabilización (aumentado)
    print("[+] Entorno listo.")

def restart_containers(env_path):
    """Reinicio rápido de contenedores entre réplicas."""
    try:
        subprocess.run(["docker-compose", "restart"], cwd=env_path, 
                       capture_output=True, timeout=60)
    except subprocess.TimeoutExpired:
        print("    [WARN] docker-compose restart colgó. Continuando...")
    time.sleep(15)

# ───────────────────────────────────────────────────────────────────────
# PARSE MD REPORT
# ───────────────────────────────────────────────────────────────────────
def parse_md_report(md_path):
    """Parsea un reporte consolidado .md y devuelve métricas estructuradas."""
    if not os.path.isfile(md_path):
        return None
    with open(md_path, "r", encoding="utf-8", errors="replace") as f:
        txt = f.read()
    
    fecha = datetime.fromtimestamp(os.path.getmtime(md_path)).strftime("%d/%m/%Y")
    hora = datetime.fromtimestamp(os.path.getmtime(md_path)).strftime("%H:%M")
    
    # Métricas K6 por Run (compatible con 4 o 5 columnas: la 5ta es Tasa de Bloqueo)
    run_metrics = {}
    for m in re.finditer(
        r'\|\s*Run\s*(\d+)\s*\|\s*([\d.]+)\s*\|\s*([\d.]+)\s*ms\s*\|\s*([\d.]+)%(?:\s*\|\s*([\d.]+)%)?',
        txt
    ):
        run_metrics[int(m.group(1))] = {
            "peticiones": float(m.group(2)),
            "latencia": float(m.group(3)),
            "fallos": float(m.group(4)),
            "bloqueo_control": float(m.group(5)) if m.group(5) else 0.0,
        }
    
    # Aislar secciones por marcadores para que los regex no crucen fronteras
    def section_between(text, start_marker, end_marker=None):
        i = text.find(start_marker)
        if i < 0: return ""
        if end_marker:
            j = text.find(end_marker, i + len(start_marker))
            return text[i:j] if j > 0 else text[i:]
        return text[i:]

    sec_cpu_pico  = section_between(txt, "PICOS DE CPU", "PICOS DE MEMORIA")
    sec_ram_pico  = section_between(txt, "PICOS DE MEMORIA", "PROMEDIO DE CPU")
    sec_cpu_mean  = section_between(txt, "PROMEDIO DE CPU", "PROMEDIO DE MEMORIA")
    sec_ram_mean  = section_between(txt, "PROMEDIO DE MEMORIA")

    # CPU pico por contenedor y run (soporta n arbitrario; excluye Promedio Pico final)
    cpu = {}
    for c in CONTAINERS:
        pat = r'\|\s*\*\*[^*]*' + re.escape(c) + r'[^*]*\*\*\s*\|(.*?)\*\*[\d.]+%\*\*'
        m = re.search(pat, sec_cpu_pico)
        if m:
            vals = re.findall(r'([\d.]+)%', m.group(0))
            # El ultimo valor es el "Promedio Pico" en negritas: se descarta
            if len(vals) >= 2:
                n_runs = len(vals) - 1
                cpu[c] = {i+1: float(vals[i]) for i in range(n_runs)}

    # CPU promedio por contenedor y run (ecuacion 4 tesis)
    cpu_mean = {}
    for c in CONTAINERS:
        pat = r'\|\s*\*\*[^*]*' + re.escape(c) + r'[^*]*\*\*\s*\|(.*?)\*\*[\d.]+%\*\*'
        m = re.search(pat, sec_cpu_mean)
        if m:
            vals = re.findall(r'([\d.]+)%', m.group(0))
            if len(vals) >= 2:
                n_runs = len(vals) - 1
                cpu_mean[c] = {i+1: float(vals[i]) for i in range(n_runs)}

    # RAM pico por contenedor y run (no incluye promedio final en la seccion 3)
    ram = {}
    for line in sec_ram_pico.split('\n'):
        if "|" in line and "**" in line:
            for c in CONTAINERS:
                if c in line:
                    mibs = re.findall(r'\(([\d.]+)\s*MiB\)', line)
                    if mibs:
                        ram[c] = {i+1: float(mibs[i]) for i in range(len(mibs))}

    # RAM promedio por contenedor y run (excluye Promedio Pico final en la seccion 5)
    ram_mean = {}
    for line in sec_ram_mean.split('\n'):
        if "|" in line and "**" in line:
            for c in CONTAINERS:
                if c in line:
                    vals = re.findall(r'([\d.]+)\s*MiB', line)
                    if len(vals) >= 2:
                        n_runs = len(vals) - 1
                        ram_mean[c] = {i+1: float(vals[i]) for i in range(n_runs)}
    
    # Promedios
    txt_clean = txt.replace('*', '')
    avg = {"peticiones": 0.0, "latencia": 0.0, "fallos": 0.0}
    m_avg = re.search(r'PROMEDIO[^|]*\|\s*([\d.]+)\s*\|\s*([\d.]+)\s*ms\s*\|\s*([\d.]+)%', txt_clean)
    if m_avg:
        avg["peticiones"] = float(m_avg.group(1))
        avg["latencia"] = float(m_avg.group(2))
        avg["fallos"] = float(m_avg.group(3))
    
    gw_vals = list(cpu.get("api-gateway", {}).values())
    avg["cpu_gw"] = sum(gw_vals) / len(gw_vals) if gw_vals else 0.0
    gw_mean_vals = list(cpu_mean.get("api-gateway", {}).values())
    avg["cpu_gw_mean"] = sum(gw_mean_vals) / len(gw_mean_vals) if gw_mean_vals else 0.0
    avg["throughput_rps"] = round(avg["peticiones"] / 60.0, 3)  # ecuacion 6 tesis
    # Ecuacion 5: D=0 si sin servicio (0 peticiones) o si fallos >= 50%
    avg["disponibilidad"] = 0 if (avg["peticiones"] == 0 or avg["fallos"] >= 50.0) else 1
    
    # Construir runs
    runs = []
    T_FASE = 60.0  # ventana de observacion por escenario (segundos)
    for n in sorted(run_metrics.keys()):
        rm = run_metrics[n]
        # Throughput (ecuacion 6): TP = N_total / T_fase (req/s)
        tp = rm["peticiones"] / T_FASE if T_FASE else 0.0
        # Disponibilidad (ecuacion 5): D = 1 si el gateway procesa, 0 si colapsa
        # Criterio operativo: D=0 si k6 no obtuvo respuestas (peticiones=0, timeout
        # o cuelgue) o si la tasa de fallos reales (5xx/timeout) alcanza el 50%.
        disp = 0 if (rm["peticiones"] == 0 or rm["fallos"] >= 50.0) else 1
        runs.append({
            "idx": n, "peticiones": rm["peticiones"], "latencia": rm["latencia"],
            "fallos": rm["fallos"], "fecha": fecha, "hora": hora,
            "throughput_rps": round(tp, 3),
            "disponibilidad": disp,
            "bloqueo_control": rm.get("bloqueo_control", 0.0),
            "cpu_gw": cpu.get("api-gateway", {}).get(n, 0.0),
            "cpu_gw_mean": cpu_mean.get("api-gateway", {}).get(n, 0.0),
            "cpu_ms_usuarios": cpu.get("ms-usuarios", {}).get(n, 0.0),
            "cpu_ms_catalogo": cpu.get("ms-catalogo", {}).get(n, 0.0),
            "cpu_ms_resenas": cpu.get("ms-resenas", {}).get(n, 0.0),
            "cpu_ms_ordenes": cpu.get("ms-ordenes", {}).get(n, 0.0),
            "cpu_mongo": cpu.get("mongo-db", {}).get(n, 0.0),
            "cpu_pg": cpu.get("postgres-db", {}).get(n, 0.0),
            "ram_gw": ram.get("api-gateway", {}).get(n, 0.0),
            "ram_gw_mean": ram_mean.get("api-gateway", {}).get(n, 0.0),
            "ram_ms_usuarios": ram.get("ms-usuarios", {}).get(n, 0.0),
            "ram_ms_catalogo": ram.get("ms-catalogo", {}).get(n, 0.0),
            "ram_ms_resenas": ram.get("ms-resenas", {}).get(n, 0.0),
            "ram_ms_ordenes": ram.get("ms-ordenes", {}).get(n, 0.0),
            "ram_mongo": ram.get("mongo-db", {}).get(n, 0.0),
            "ram_pg": ram.get("postgres-db", {}).get(n, 0.0),
        })
    
    return {"fecha": fecha, "hora": hora, "runs": runs, "avg": avg}

# ───────────────────────────────────────────────────────────────────────
# EXCEL EXPORT
# ───────────────────────────────────────────────────────────────────────
def export_uc01_excel(full_data, avg_data, vus_completed, out_path, breaking_point=None):
    """Genera el Excel comparativo para UC-01."""
    if not HAS_EXCEL:
        print("[WARN] openpyxl no instalado. No se genera Excel.")
        return
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UC01 Introspeccion"
    ws.sheet_view.showGridLines = False
    
    thin = Side(style="thin", color="000000")
    bdr = Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ttl_font = Font(bold=True, size=13, color="FFFFFF")
    hdr_font = Font(bold=True, color="FFFFFF", size=9)
    data_font = Font(color="000000", name="Calibri", size=9)
    fill_title = PatternFill("solid", fgColor="2F5597")
    fill_hdr = PatternFill("solid", fgColor="2F5597")
    fill_vuln = PatternFill("solid", fgColor="FFF2CC")
    fill_prot = PatternFill("solid", fgColor="D9E1F2")
    fill_break = PatternFill("solid", fgColor="FF6B6B")
    
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    bp_str = f" | Punto de Quiebre: {breaking_point} VUs" if breaking_point else ""
    ws.merge_cells("A1:AD1")
    c = ws["A1"]
    c.value = f"UC-01 Introspección - Fibonacci Extendido{bp_str} - {now_str}"
    c.font = ttl_font; c.fill = fill_title; c.alignment = center

    HEADERS = [
        "Caso de Uso", "Entorno", "Tratamiento", "Carga (VUs)", "Replica",
        "Fecha", "Hora",
        "Peticiones Totales", "Throughput (req/s)", "Latencia (ms)",
        "Fallos (%)", "Bloqueo Control (%)", "Disponibilidad (D)",
        "GW CPU Pico (%)", "GW CPU Prom (%)",
        "ms-usuarios CPU (%)", "ms-catalogo CPU (%)", "ms-resenas CPU (%)",
        "ms-ordenes CPU (%)", "Mongo CPU (%)", "Postgres CPU (%)",
        "GW RAM Pico (MiB)", "GW RAM Prom (MiB)",
        "ms-usuarios RAM (MiB)", "ms-catalogo RAM (MiB)", "ms-resenas RAM (MiB)",
        "ms-ordenes RAM (MiB)", "Mongo RAM (MiB)", "Postgres RAM (MiB)",
        "Tipo"
    ]
    
    for ci, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = hdr_font; cell.fill = fill_hdr; cell.alignment = center; cell.border = bdr
    
    row_idx = 4
    for vus in vus_completed:
        for env_name in ["Vulnerable", "Protegido"]:
            is_vuln = (env_name == "Vulnerable")
            fill = fill_vuln if is_vuln else fill_prot
            tratam = "Linea Base Vulnerable" if is_vuln else "Hardening ISO 27001"
            
            fd = full_data.get(env_name, {}).get(vus, {})
            runs = fd.get("runs", [])
            
            if not runs:
                # Fila vacía indicando sin datos
                cell = ws.cell(row=row_idx, column=1, value="UC-01")
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                cell = ws.cell(row=row_idx, column=2, value=env_name)
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                cell = ws.cell(row=row_idx, column=4, value=vus)
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                cell = ws.cell(row=row_idx, column=len(HEADERS), value="Sin Datos")
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                row_idx += 1
                continue

            for run in runs:
                is_break = (breaking_point and vus >= breaking_point and is_vuln)
                row_fill = fill_break if is_break else fill
                peticiones = run['peticiones']
                # Ecuacion 6: TP = N/T_fase (req/s)
                tp_rps = round(run.get('throughput_rps', peticiones / 60.0), 3)
                # Ecuacion 5: D binaria (1 si procesa, 0 si colapsa por fallos >= 50%)
                disp = run.get('disponibilidad', 0 if run['fallos'] >= 50.0 else 1)
                # Bloqueo del control: viene de la metrica REAL 'iso_tasa_bloqueadas'
                # emitida por los scripts K6 modificados. Si el reporte es previo a esa
                # actualizacion, el campo queda en 0 (n/m = no medido en esa corrida).
                bloqueo = run.get('bloqueo_control', 0.0)
                vals = [
                    "UC-01", env_name, tratam, vus, f"R{run['idx']}",
                    run['fecha'], run['hora'],
                    round(peticiones, 2), tp_rps, round(run['latencia'], 2),
                    round(run['fallos'], 2), round(bloqueo, 2), disp,
                    round(run.get('cpu_gw', 0), 2), round(run.get('cpu_gw_mean', 0), 2),
                    round(run.get('cpu_ms_usuarios', 0), 2),
                    round(run.get('cpu_ms_catalogo', 0), 2),
                    round(run.get('cpu_ms_resenas', 0), 2),
                    round(run.get('cpu_ms_ordenes', 0), 2),
                    round(run.get('cpu_mongo', 0), 2),
                    round(run.get('cpu_pg', 0), 2),
                    round(run.get('ram_gw', 0), 2), round(run.get('ram_gw_mean', 0), 2),
                    round(run.get('ram_ms_usuarios', 0), 2),
                    round(run.get('ram_ms_catalogo', 0), 2),
                    round(run.get('ram_ms_resenas', 0), 2),
                    round(run.get('ram_ms_ordenes', 0), 2),
                    round(run.get('ram_mongo', 0), 2),
                    round(run.get('ram_pg', 0), 2),
                    "Punto de Quiebre" if is_break else "Replica"
                ]
                for ci, v in enumerate(vals, 1):
                    cell = ws.cell(row=row_idx, column=ci, value=v)
                    cell.fill = row_fill; cell.alignment = center
                    cell.border = bdr; cell.font = data_font
                row_idx += 1
    
    # Auto-ajuste de columnas
    col_widths = [10, 14, 22, 10, 8, 12, 8,
                  16, 16, 14, 12, 16, 14,
                  15, 15,
                  16, 16, 15, 15, 13, 15,
                  16, 16,
                  18, 18, 17, 17, 15, 15,
                  16]
    for ci, w in enumerate(col_widths, 1):
        if ci <= len(HEADERS):
            ws.column_dimensions[get_column_letter(ci)].width = w
    
    ws.freeze_panes = "A4"
    wb.save(out_path)
    print(f"    [EXCEL] Guardado: {out_path}")

# ───────────────────────────────────────────────────────────────────────
# GRAFICOS
# ───────────────────────────────────────────────────────────────────────
def generate_uc01_graphs(avg_data, vus_completed, breaking_point=None):
    """Genera gráficos comparativos Vulnerable vs Protegido."""
    BG = "#ffffff"; BG2 = "#ffffff"; BORD = "#d0d7de"
    FG = "#1f2328"; MUTED = "#57606a"

    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 9,
        "figure.facecolor": BG2, "axes.facecolor": BG,
        "axes.edgecolor": BORD, "axes.labelcolor": FG,
        "xtick.color": MUTED, "ytick.color": MUTED,
        "grid.color": BORD, "grid.alpha": 0.6, "axes.grid": True,
        "grid.linestyle": "--", "legend.facecolor": "#ffffff",
        "legend.edgecolor": BORD, "legend.labelcolor": FG, "text.color": FG,
    })
    
    metrics = [
        ("Throughput (Peticiones Totales)", "peticiones", "req"),
        ("Latencia Media (ms)", "latencia", "ms"),
        ("Tasa de Fallos (%)", "fallos", "%"),
        ("CPU API Gateway (%)", "cpu_gw", "%"),
    ]
    
    for metric_name, metric_key, unit in metrics:
        fig, ax = plt.subplots(figsize=(12, 6))
        fig.patch.set_facecolor(BG2)
        
        for env_name, color, marker in [("Vulnerable", "#f85149", "o"), ("Protegido", "#3fb950", "s")]:
            env_avg = avg_data.get(env_name, {})
            vus_vals = []
            met_vals = []
            for v in vus_completed:
                if v in env_avg:
                    vus_vals.append(v)
                    met_vals.append(env_avg[v].get(metric_key, 0))
            
            if vus_vals:
                ax.plot(range(len(vus_vals)), met_vals, marker=marker, color=color,
                        lw=2.2, ms=8, label=env_name, zorder=3)
                for i, (v, val) in enumerate(zip(vus_vals, met_vals)):
                    off = 12 if env_name == "Vulnerable" else -15
                    ax.annotate(f"{val:.1f}", (i, val), textcoords="offset points",
                                xytext=(0, off), ha="center", fontsize=7.5, color=color)
        
        # Línea vertical de punto de quiebre
        if breaking_point and breaking_point in vus_completed:
            bp_idx = vus_completed.index(breaking_point)
            ax.axvline(x=bp_idx, color="#d29922", linestyle=":", lw=2, alpha=0.8,
                       label=f"Punto de Quiebre ({breaking_point} VUs)")
        
        ax.set_xticks(range(len(vus_completed)))
        ax.set_xticklabels([str(v) for v in vus_completed], rotation=45)
        ax.set_xlabel("Usuarios Virtuales (Fibonacci)")
        ax.set_ylabel(f"{metric_name} ({unit})")
        ax.set_title(f"UC-01 Introspección — {metric_name}", fontsize=13, pad=12,
                      fontweight="bold", color=FG)
        ax.legend(loc="upper left", fontsize=9)
        
        plt.tight_layout()
        safe_name = metric_name.replace(" ", "_").replace("(", "").replace(")", "").replace("/", "_").replace("%", "pct")
        out_png = os.path.join(OUT_DIR, f"UC01_Introspeccion_{safe_name}.png")
        fig.savefig(out_png, dpi=150, bbox_inches="tight")
        plt.close(fig)
        print(f"    [GRAFICO] {out_png}")

# ───────────────────────────────────────────────────────────────────────
# RUNNER PRINCIPAL
# ───────────────────────────────────────────────────────────────────────
def run_uc01():
    print("=" * 70)
    print("  UC-01: INTROSPECCIÓN — BÚSQUEDA DE PUNTO DE QUIEBRE (FIBONACCI)")
    print("=" * 70)
    
    full_data = {"Vulnerable": {}, "Protegido": {}}
    avg_data = {"Vulnerable": {}, "Protegido": {}}
    vus_completed = []
    breaking_point = None
    
    for vus in FIBONACCI_VUS:
        print(f"\n{'='*60}")
        print(f"  VUs = {vus}  (Fibonacci)")
        print(f"{'='*60}")
        
        for env_name, env_path in ENVS.items():
            print(f"\n  --- [{env_name}] ---")
            md_path = os.path.join(env_path, "load_tests", "UC01_introspeccion",
                                   "resultados", f"reporte_consolidado_{RUNS}runs_vus{vus}.md")

            if os.path.exists(md_path):
                print(f"    [SKIP] Reporte ya existe para VUs={vus}, reutilizando resultados previos...")
            else:
                restart_env_full(env_path)

                run_script = os.path.join(env_path, "load_tests", "run_multiple_experiments.py")

                print(f"    Ejecutando {RUNS} réplicas con {vus} VUs...")
                cmd = [sys.executable, run_script, "--test-script", UC_SCRIPT,
                       "--vus", str(vus), "--runs", str(RUNS)]

                try:
                    p = subprocess.Popen(cmd, cwd=env_path,
                                         stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                         text=True)
                    for line in p.stdout:
                        print(f"      {line.rstrip()}")
                    p.wait(timeout=600)  # 10 min max por bloque
                    if p.returncode != 0:
                        print(f"    [ERROR] El orquestador de K6 falló con código {p.returncode}. Continuando con siguiente...")
                        continue
                except subprocess.TimeoutExpired:
                    print(f"    [TIMEOUT] La prueba con {vus} VUs tardó más de 10 min. Matando...")
                    p.kill()
                    p.wait()
                except Exception as e:
                    print(f"    [ERROR] {e}")

            # Leer el reporte generado (recien creado o reutilizado)
            if os.path.exists(md_path):
                res = parse_md_report(md_path)
                if res:
                    full_data[env_name][vus] = res
                    avg_data[env_name][vus] = res["avg"]
                    print(f"    [OK] Peticiones: {res['avg']['peticiones']:.0f} | "
                          f"Latencia: {res['avg']['latencia']:.2f} ms | "
                          f"Fallos: {res['avg']['fallos']:.2f}%")
            else:
                print(f"    [ERROR] No se generó el reporte: {md_path}")
        
        vus_completed.append(vus)
        
        # Verificar punto de quiebre en el Vulnerable
        vuln_avg = avg_data.get("Vulnerable", {}).get(vus, {})
        fail_rate = vuln_avg.get("fallos", 0.0)
        
        if fail_rate >= FAIL_THRESHOLD and breaking_point is None:
            breaking_point = vus
            print(f"\n  *** POSIBLE PUNTO DE QUIEBRE DETECTADO: {vus} VUs (Fallos: {fail_rate:.2f}%) ***")
            print("  Se evaluara el siguiente nivel para confirmar...")
        elif fail_rate < FAIL_THRESHOLD and breaking_point is not None:
            print(f"\n  [FALSO POSITIVO] El nivel {vus} se recuperó con solo {fail_rate:.2f}% de fallos.")
            print(f"  Descartando el quiebre anterior en {breaking_point} VUs y continuando...")
            breaking_point = None
        
        # Generar Excel y gráficos progresivos después de cada nivel
        print(f"\n  [+] Generando reportes parciales (VUs completados: {vus_completed})...")
        excel_path = os.path.join(OUT_DIR, "UC01_Introspeccion_Fibonacci_Comparativa.xlsx")
        try:
            export_uc01_excel(full_data, avg_data, vus_completed, excel_path, breaking_point)
        except PermissionError:
            print(f"  [ERROR] Permiso denegado para guardar Excel. Asegúrate de cerrarlo en Microsoft Excel.")
        except Exception as e:
            print(f"  [ERROR] Fallo al generar Excel: {e}")
            
        generate_uc01_graphs(avg_data, vus_completed, breaking_point)
        
        # Si ya encontramos el punto de quiebre, hacer un nivel más para confirmar y parar
        if breaking_point and vus > breaking_point and fail_rate >= FAIL_THRESHOLD:
            print(f"\n  [STOP] Confirmado colapso total. "
                  f"Nivel {vus} confirma la tendencia ({fail_rate:.2f}% fallos). Deteniendo.")
            break
    
    # Apagar entornos
    print("\n[+] Apagando entornos Docker...")
    for env_name, env_path in ENVS.items():
        try:
            subprocess.run(["docker-compose", "down"], cwd=env_path, 
                          capture_output=True, timeout=60)
        except:
            pass
    
    # Resumen final
    print("\n" + "=" * 70)
    print("  RESUMEN FINAL UC-01")
    print("=" * 70)
    if breaking_point:
        print(f"  PUNTO DE QUIEBRE: {breaking_point} VUs")
        print(f"  El servidor Vulnerable colapsa a partir de {breaking_point} usuarios concurrentes.")
    else:
        print(f"  No se encontró punto de quiebre con la secuencia probada: {vus_completed}")
        print(f"  El servidor resistió hasta {vus_completed[-1]} VUs.")
    
    print(f"\n  Archivos generados en: {OUT_DIR}")
    print(f"  VUs evaluados: {vus_completed}")
    print("=" * 70)
    
    return breaking_point, vus_completed

if __name__ == "__main__":
    start = datetime.now()
    try:
        bp, vus_done = run_uc01()
    except KeyboardInterrupt:
        print("\n[CANCELADO] Prueba interrumpida por el usuario.")
    except Exception as e:
        print(f"\n[ERROR FATAL] {e}")
        import traceback; traceback.print_exc()
    finally:
        end = datetime.now()
        print(f"\n[{end}] Duración total: {end - start}")
