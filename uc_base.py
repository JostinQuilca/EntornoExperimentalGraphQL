# -*- coding: utf-8 -*-
"""
Modulo base para runners UC-02 al UC-05.
Comparte con uc01_runner.py: reinicio Docker, parseo de reportes, exportacion
Excel y generacion de graficas. La diferencia clave es que aqui la variable
independiente NO es VUs sino un parametro de escenario (profundidad, ciclos,
alias, fragmentos) inyectado por variable de entorno.
"""
import os, sys, time, subprocess, re, urllib.request, urllib.error
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

ENVS = {
    "Vulnerable": os.path.join(BASE_DIR, "E-commerce"),
    "Protegido":  os.path.join(BASE_DIR, "E-commerce-controles"),
}

CONTAINERS = ['api-gateway', 'mongo-db', 'ms-catalogo', 'ms-ordenes',
              'ms-resenas', 'ms-usuarios', 'postgres-db']


# ─────────────────────────────────────────────────────────────
# DOCKER HELPERS
# ─────────────────────────────────────────────────────────────
def restart_env_full(env_path):
    print(f"\n[+] Full Restart de {env_path}...")
    subprocess.run(["docker-compose", "down", "--remove-orphans"], cwd=ENVS["Vulnerable"],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
    subprocess.run(["docker-compose", "down", "--remove-orphans"], cwd=ENVS["Protegido"],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
    result = None
    for _retry in range(3):
        time.sleep(5)
        result = subprocess.run(["docker-compose", "up", "-d"], cwd=env_path,
                                stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=600)
        if result.returncode == 0:
            break
        print(f"[WARN] docker-compose up fallo (intento {_retry+1}/3). Reintentando...")
        subprocess.run(["docker-compose", "down"], cwd=env_path,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
        time.sleep(3)
    else:
        print(f"[ERROR] docker-compose up fallo tras 3 intentos: {result.stderr.decode('utf-8', errors='ignore') if result else ''}")
        sys.exit(1)

    print("[HEALTHCHECK] Verificando gateway...")
    for _attempt in range(30):
        time.sleep(2)
        try:
            req = urllib.request.Request("http://localhost:4000/graphql",
                data=b'{"query":"{ __typename }"}',
                headers={"Content-Type": "application/json"}, method="POST")
            resp = urllib.request.urlopen(req, timeout=5)
            print(f"[OK] Gateway respondio (status={resp.status})")
            break
        except urllib.error.HTTPError as e:
            if e.code in (400, 403):
                print(f"[OK] Gateway respondio (status={e.code})")
                break
        except Exception:
            pass
    else:
        print("[WARN] Gateway no respondio tras 60s.")
    time.sleep(8)


# ─────────────────────────────────────────────────────────────
# PARSE MD REPORT (mismo formato que uc01_runner.py)
# ─────────────────────────────────────────────────────────────
def parse_md_report(md_path):
    if not os.path.isfile(md_path): return None
    with open(md_path, "r", encoding="utf-8", errors="replace") as f:
        txt = f.read()

    fecha = datetime.fromtimestamp(os.path.getmtime(md_path)).strftime("%d/%m/%Y")
    hora = datetime.fromtimestamp(os.path.getmtime(md_path)).strftime("%H:%M")

    run_metrics = {}
    for m in re.finditer(
        r'\|\s*Run\s*(\d+)\s*\|\s*([\d.]+)\s*\|\s*([\d.]+)\s*ms\s*\|\s*([\d.]+)%(?:\s*\|\s*([\d.]+)%)?',
        txt
    ):
        run_metrics[int(m.group(1))] = {
            "peticiones": float(m.group(2)),
            "latencia":   float(m.group(3)),
            "fallos":     float(m.group(4)),
            "bloqueo_control": float(m.group(5)) if m.group(5) else 0.0,
        }

    def section(text, start, end=None):
        i = text.find(start)
        if i < 0: return ""
        if end:
            j = text.find(end, i + len(start))
            return text[i:j] if j > 0 else text[i:]
        return text[i:]

    sec_cpu_pico  = section(txt, "PICOS DE CPU", "PICOS DE MEMORIA")
    sec_ram_pico  = section(txt, "PICOS DE MEMORIA", "PROMEDIO DE CPU")
    sec_cpu_mean  = section(txt, "PROMEDIO DE CPU", "PROMEDIO DE MEMORIA")
    sec_ram_mean  = section(txt, "PROMEDIO DE MEMORIA")

    def parse_container_table(txt_section, unit_regex):
        # Soporta n arbitrario. Regex termina antes del "Promedio Pico" (**X.XX%**)
        result = {}
        for c in CONTAINERS:
            pat = r'\|\s*\*\*[^*]*' + re.escape(c) + r'[^*]*\*\*\s*\|(.*?)\*\*[\d.]+[%A-Za-z]+\*\*'
            m = re.search(pat, txt_section)
            if not m:
                # Fallback: si no hay promedio final, tomar toda la linea
                pat_fb = r'\|\s*\*\*[^*]*' + re.escape(c) + r'[^*]*\*\*\s*\|([^\n]*)'
                m = re.search(pat_fb, txt_section)
            if m:
                vals = re.findall(unit_regex, m.group(1))
                if vals:
                    result[c] = {i+1: float(vals[i]) for i in range(len(vals))}
        return result

    cpu       = parse_container_table(sec_cpu_pico, r'([\d.]+)%')
    cpu_mean  = parse_container_table(sec_cpu_mean, r'([\d.]+)%')

    ram = {}
    for line in sec_ram_pico.split('\n'):
        if "|" in line and "**" in line:
            for c in CONTAINERS:
                if c in line:
                    mibs = re.findall(r'\(([\d.]+)\s*MiB\)', line)
                    if mibs:
                        ram[c] = {i+1: float(mibs[i]) for i in range(len(mibs))}

    ram_mean = {}
    for line in sec_ram_mean.split('\n'):
        if "|" in line and "**" in line:
            for c in CONTAINERS:
                if c in line:
                    # Excluir el "Promedio Pico" final si esta en negritas
                    line_no_avg = re.sub(r'\*\*[\d.]+\s*MiB\*\*.*$', '', line)
                    vals = re.findall(r'([\d.]+)\s*MiB', line_no_avg)
                    if vals:
                        ram_mean[c] = {i+1: float(vals[i]) for i in range(len(vals))}

    avg = {"peticiones": 0.0, "latencia": 0.0, "fallos": 0.0}
    m_avg = re.search(r'PROMEDIO[^|]*\|\s*([\d.]+)\s*\|\s*([\d.]+)\s*ms\s*\|\s*([\d.]+)%', txt.replace('*', ''))
    if m_avg:
        avg["peticiones"] = float(m_avg.group(1))
        avg["latencia"]   = float(m_avg.group(2))
        avg["fallos"]     = float(m_avg.group(3))
    gw_vals = list(cpu.get("api-gateway", {}).values())
    avg["cpu_gw"] = sum(gw_vals) / len(gw_vals) if gw_vals else 0.0
    gw_mean_vals = list(cpu_mean.get("api-gateway", {}).values())
    avg["cpu_gw_mean"] = sum(gw_mean_vals) / len(gw_mean_vals) if gw_mean_vals else 0.0
    avg["throughput_rps"] = round(avg["peticiones"] / 60.0, 3)
    avg["disponibilidad"] = 0 if (avg["peticiones"] == 0 or avg["fallos"] >= 50.0) else 1

    runs = []
    for n in sorted(run_metrics.keys()):
        rm = run_metrics[n]
        disp = 0 if (rm["peticiones"] == 0 or rm["fallos"] >= 50.0) else 1
        runs.append({
            "idx": n, "peticiones": rm["peticiones"], "latencia": rm["latencia"],
            "fallos": rm["fallos"], "fecha": fecha, "hora": hora,
            "throughput_rps": round(rm["peticiones"] / 60.0, 3),
            "disponibilidad": disp,
            "bloqueo_control": rm.get("bloqueo_control", 0.0),
            "cpu_gw": cpu.get("api-gateway", {}).get(n, 0.0),
            "cpu_gw_mean": cpu_mean.get("api-gateway", {}).get(n, 0.0),
            "cpu_ms_usuarios": cpu.get("ms-usuarios", {}).get(n, 0.0),
            "cpu_ms_catalogo": cpu.get("ms-catalogo", {}).get(n, 0.0),
            "cpu_ms_resenas": cpu.get("ms-resenas", {}).get(n, 0.0),
            "cpu_ms_ordenes": cpu.get("ms-ordenes", {}).get(n, 0.0),
            "cpu_mongo": cpu.get("mongo-db", {}).get(n, 0.0),
            "cpu_pg": cpu.get("postgres-db", {}).get(n, 0.0),
            "ram_gw": ram.get("api-gateway", {}).get(n, 0.0),
            "ram_gw_mean": ram_mean.get("api-gateway", {}).get(n, 0.0),
            "ram_ms_usuarios": ram.get("ms-usuarios", {}).get(n, 0.0),
            "ram_ms_catalogo": ram.get("ms-catalogo", {}).get(n, 0.0),
            "ram_ms_resenas": ram.get("ms-resenas", {}).get(n, 0.0),
            "ram_ms_ordenes": ram.get("ms-ordenes", {}).get(n, 0.0),
            "ram_mongo": ram.get("mongo-db", {}).get(n, 0.0),
            "ram_pg": ram.get("postgres-db", {}).get(n, 0.0),
        })

    return {"fecha": fecha, "hora": hora, "runs": runs, "avg": avg}


# ─────────────────────────────────────────────────────────────
# EXCEL EXPORT genérico
# ─────────────────────────────────────────────────────────────
def export_excel(uc_id, var_label, full_data, levels_completed, out_path, vus_fijo):
    if not HAS_EXCEL:
        print("[WARN] openpyxl no instalado.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = uc_id.replace('-', '')
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color="000000")
    bdr = Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ttl_font = Font(bold=True, size=13, color="FFFFFF")
    hdr_font = Font(bold=True, color="FFFFFF", size=9)
    data_font = Font(color="000000", name="Calibri", size=9)
    fill_title = PatternFill("solid", fgColor="2F5597")
    fill_hdr = PatternFill("solid", fgColor="2F5597")
    fill_vuln = PatternFill("solid", fgColor="FFF2CC")
    fill_prot = PatternFill("solid", fgColor="D9E1F2")

    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A1:AE1")
    c = ws["A1"]
    c.value = f"{uc_id} - Comparativa Vulnerable vs Protegido - VUs={vus_fijo} - {now_str}"
    c.font = ttl_font; c.fill = fill_title; c.alignment = center

    HEADERS = [
        "Caso de Uso", "Entorno", "Tratamiento", var_label, "VUs", "Replica",
        "Fecha", "Hora",
        "Peticiones Totales", "Throughput (req/s)", "Latencia (ms)",
        "Fallos (%)", "Bloqueo Control (%)", "Disponibilidad (D)",
        "GW CPU Pico (%)", "GW CPU Prom (%)",
        "ms-usuarios CPU (%)", "ms-catalogo CPU (%)", "ms-resenas CPU (%)",
        "ms-ordenes CPU (%)", "Mongo CPU (%)", "Postgres CPU (%)",
        "GW RAM Pico (MiB)", "GW RAM Prom (MiB)",
        "ms-usuarios RAM (MiB)", "ms-catalogo RAM (MiB)", "ms-resenas RAM (MiB)",
        "ms-ordenes RAM (MiB)", "Mongo RAM (MiB)", "Postgres RAM (MiB)",
        "Tipo"
    ]

    for ci, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = hdr_font; cell.fill = fill_hdr; cell.alignment = center; cell.border = bdr

    row_idx = 4
    for level in levels_completed:
        for env_name in ["Vulnerable", "Protegido"]:
            is_vuln = (env_name == "Vulnerable")
            fill = fill_vuln if is_vuln else fill_prot
            tratam = "Linea Base Vulnerable" if is_vuln else "Hardening ISO 27001"
            fd = full_data.get(env_name, {}).get(level, {})
            runs = fd.get("runs", [])
            if not runs:
                cell = ws.cell(row=row_idx, column=1, value=uc_id)
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                cell = ws.cell(row=row_idx, column=2, value=env_name)
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                cell = ws.cell(row=row_idx, column=4, value=level)
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                cell = ws.cell(row=row_idx, column=len(HEADERS), value="Sin Datos")
                cell.fill = fill; cell.alignment = center; cell.border = bdr
                row_idx += 1
                continue
            for run in runs:
                # Bloqueo del control: viene de la metrica REAL 'iso_tasa_bloqueadas'
                # emitida por los scripts K6 modificados. Reportes previos = 0 (n/m).
                bloqueo = run.get('bloqueo_control', 0.0)
                vals = [
                    uc_id, env_name, tratam, level, vus_fijo, f"R{run['idx']}",
                    run['fecha'], run['hora'],
                    round(run['peticiones'], 2), run.get('throughput_rps', 0),
                    round(run['latencia'], 2),
                    round(run['fallos'], 2), round(bloqueo, 2),
                    run.get('disponibilidad', 1),
                    round(run.get('cpu_gw', 0), 2), round(run.get('cpu_gw_mean', 0), 2),
                    round(run.get('cpu_ms_usuarios', 0), 2),
                    round(run.get('cpu_ms_catalogo', 0), 2),
                    round(run.get('cpu_ms_resenas', 0), 2),
                    round(run.get('cpu_ms_ordenes', 0), 2),
                    round(run.get('cpu_mongo', 0), 2),
                    round(run.get('cpu_pg', 0), 2),
                    round(run.get('ram_gw', 0), 2), round(run.get('ram_gw_mean', 0), 2),
                    round(run.get('ram_ms_usuarios', 0), 2),
                    round(run.get('ram_ms_catalogo', 0), 2),
                    round(run.get('ram_ms_resenas', 0), 2),
                    round(run.get('ram_ms_ordenes', 0), 2),
                    round(run.get('ram_mongo', 0), 2),
                    round(run.get('ram_pg', 0), 2),
                    "Replica"
                ]
                for ci, v in enumerate(vals, 1):
                    cell = ws.cell(row=row_idx, column=ci, value=v)
                    cell.fill = fill; cell.alignment = center
                    cell.border = bdr; cell.font = data_font
                row_idx += 1

    for ci in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 15
    ws.freeze_panes = "A4"
    wb.save(out_path)
    print(f"    [EXCEL] Guardado: {out_path}")


# ─────────────────────────────────────────────────────────────
# GRAFICAS genéricas (tema claro, consistente con UC-01)
# ─────────────────────────────────────────────────────────────
def generate_graphs(uc_id, var_label, avg_data, levels_completed, out_dir):
    BG = "#ffffff"; BORD = "#d0d7de"; FG = "#1f2328"; MUTED = "#57606a"
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 9,
        "figure.facecolor": BG, "axes.facecolor": BG,
        "axes.edgecolor": BORD, "axes.labelcolor": FG,
        "xtick.color": MUTED, "ytick.color": MUTED,
        "grid.color": BORD, "grid.alpha": 0.6, "axes.grid": True,
        "grid.linestyle": "--", "legend.facecolor": BG,
        "legend.edgecolor": BORD, "legend.labelcolor": FG, "text.color": FG,
    })

    metrics = [
        ("Throughput (peticiones)", "peticiones", "req"),
        ("Latencia Media (ms)", "latencia", "ms"),
        ("Tasa de Fallos (%)", "fallos", "%"),
        ("CPU API Gateway Promedio (%)", "cpu_gw_mean", "%"),
    ]

    for metric_name, metric_key, unit in metrics:
        fig, ax = plt.subplots(figsize=(12, 6))
        for env_name, color, marker in [("Vulnerable", "#e11d48", "o"), ("Protegido", "#059669", "s")]:
            env_avg = avg_data.get(env_name, {})
            xs, ys = [], []
            for lv in levels_completed:
                if lv in env_avg:
                    xs.append(lv)
                    ys.append(env_avg[lv].get(metric_key, 0))
            if xs:
                ax.plot(range(len(xs)), ys, marker=marker, color=color,
                        lw=2.2, ms=8, label=env_name, zorder=3)
                for i, (v, val) in enumerate(zip(xs, ys)):
                    off = 12 if env_name == "Vulnerable" else -15
                    ax.annotate(f"{val:.1f}", (i, val), textcoords="offset points",
                                xytext=(0, off), ha="center", fontsize=7.5, color=color)

        ax.set_xticks(range(len(levels_completed)))
        ax.set_xticklabels([str(v) for v in levels_completed], rotation=45)
        ax.set_xlabel(var_label)
        ax.set_ylabel(f"{metric_name} ({unit})")
        ax.set_title(f"{uc_id} — {metric_name}", fontsize=13, pad=12,
                     fontweight="bold", color=FG)
        ax.legend(loc="upper left", fontsize=9)
        plt.tight_layout()
        safe = metric_name.replace(" ", "_").replace("(", "").replace(")", "").replace("%", "pct").replace("/", "_")
        out_png = os.path.join(out_dir, f"{uc_id.replace('-','')}_{safe}.png")
        fig.savefig(out_png, dpi=150, bbox_inches="tight")
        plt.close(fig)
        print(f"    [GRAFICO] {out_png}")


# ─────────────────────────────────────────────────────────────
# RUNNER PRINCIPAL PARAMETRIZADO
# ─────────────────────────────────────────────────────────────
def run_uc_generic(uc_id, uc_folder, script_rel, var_env_name, var_label,
                   levels, vus_fixed, runs=3, out_subdir=None):
    """
    uc_id:        p.ej. "UC-02"
    uc_folder:    subcarpeta bajo load_tests p.ej. "UC02_profundidad_moderada"
    script_rel:   ruta al .js relativa a load_tests p.ej. "UC02_profundidad_moderada/uc02_profundidad_nivel.js"
    var_env_name: variable de entorno inyectada por nivel: "LEVEL" o "ALIAS_COUNT"
    var_label:    etiqueta legible del nivel: "Profundidad", "Ciclos", "Numero de alias", "Fragmentos"
    levels:       lista de valores del parametro variable
    vus_fixed:    VUs fijo en todos los niveles
    runs:         numero de replicas por nivel
    out_subdir:   subcarpeta bajo pruebas/ para salida. Default: uc_id sin guion.
    """
    if out_subdir is None:
        out_subdir = uc_id.replace('-', '_')
    out_dir = os.path.join(BASE_DIR, "pruebas", out_subdir)
    os.makedirs(out_dir, exist_ok=True)

    print("=" * 70)
    print(f"  {uc_id}: Corrida por niveles de {var_label}")
    print(f"  Niveles: {levels} | VUs fijo: {vus_fixed} | Replicas: {runs}")
    print("=" * 70)

    full_data = {"Vulnerable": {}, "Protegido": {}}
    avg_data = {"Vulnerable": {}, "Protegido": {}}
    levels_completed = []

    for level in levels:
        print(f"\n{'='*60}\n  {var_label} = {level}\n{'='*60}")
        for env_name, env_path in ENVS.items():
            print(f"\n  --- [{env_name}] ---")
            md_path = os.path.join(env_path, "load_tests", uc_folder, "resultados",
                                   f"reporte_consolidado_{runs}runs_vus{vus_fixed}_nivel{level}.md")
            if os.path.exists(md_path):
                print(f"    [SKIP] Reporte ya existe para {var_label}={level}, reutilizando...")
            else:
                restart_env_full(env_path)
                run_script = os.path.join(env_path, "load_tests", "run_multiple_experiments.py")
                print(f"    Ejecutando {runs} replicas con {var_label}={level}, VUs={vus_fixed}...")
                cmd = [sys.executable, run_script, "--test-script", script_rel,
                       "--vus", str(vus_fixed), "--runs", str(runs)]
                env_vars = os.environ.copy()
                env_vars[var_env_name] = str(level)
                env_vars["LEVEL"] = str(level)  # compat: algunos scripts leen LEVEL
                try:
                    p = subprocess.Popen(cmd, cwd=env_path,
                                         stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                         text=True, env=env_vars)
                    for line in p.stdout:
                        print(f"      {line.rstrip()}")
                    p.wait(timeout=1200)
                    if p.returncode != 0:
                        print(f"    [ERROR] Orquestador de K6 fallo con codigo {p.returncode}.")
                        continue
                except subprocess.TimeoutExpired:
                    print(f"    [TIMEOUT] Nivel {level} tardo mas de 20 min. Matando...")
                    p.kill(); p.wait()
                except Exception as e:
                    print(f"    [ERROR] {e}")

            if os.path.exists(md_path):
                res = parse_md_report(md_path)
                if res:
                    full_data[env_name][level] = res
                    avg_data[env_name][level] = res["avg"]
                    print(f"    [OK] Peticiones: {res['avg']['peticiones']:.0f} | "
                          f"Latencia: {res['avg']['latencia']:.2f} ms | "
                          f"Fallos: {res['avg']['fallos']:.2f}%")
            else:
                print(f"    [ERROR] No se genero el reporte: {md_path}")

        levels_completed.append(level)
        excel_path = os.path.join(out_dir, f"{uc_id.replace('-','')}_Comparativa.xlsx")
        try:
            export_excel(uc_id, var_label, full_data, levels_completed, excel_path, vus_fixed)
        except PermissionError:
            print(f"    [ERROR] Cerra el Excel para poder sobrescribirlo.")
        except Exception as e:
            print(f"    [ERROR EXCEL] {e}")
        generate_graphs(uc_id, var_label, avg_data, levels_completed, out_dir)

    print("\n[+] Apagando entornos Docker...")
    for env_path in ENVS.values():
        try:
            subprocess.run(["docker-compose", "down"], cwd=env_path,
                          capture_output=True, timeout=60)
        except Exception:
            pass

    print("\n" + "=" * 70)
    print(f"  RESUMEN {uc_id}")
    print("=" * 70)
    print(f"  Niveles evaluados: {levels_completed}")
    print(f"  Archivos generados en: {out_dir}")
    return levels_completed


# ═══════════════════════════════════════════════════════════════
# RUNNER MATRIZ: Nivel × VUs (para diseño factorial 2D)
# ═══════════════════════════════════════════════════════════════
def run_uc_matrix(uc_id, uc_folder, script_rel, var_env_name, var_label,
                  levels, vus_list, runs=5, out_subdir=None):
    """
    Ejecuta la matriz factorial completa: cada nivel se prueba con cada VUs.
    Reutiliza reportes existentes; sobre lo que falta ejecuta k6.
    Al final exporta un Excel con todas las combinaciones y un análisis matriz.
    """
    if out_subdir is None:
        out_subdir = uc_id.replace('-', '_')
    out_dir = os.path.join(BASE_DIR, "pruebas", out_subdir)
    os.makedirs(out_dir, exist_ok=True)

    print("=" * 70)
    print(f"  {uc_id}: MATRIZ {var_label} × VUs")
    print(f"  {var_label}: {levels}")
    print(f"  VUs: {vus_list}")
    print(f"  Total celdas: {len(levels) * len(vus_list)} × 2 entornos × {runs} réplicas")
    print("=" * 70)

    # Estructura: full_data[env][(level, vus)] = res
    full_data = {"Vulnerable": {}, "Protegido": {}}
    combos_done = []

    for level in levels:
        for vus in vus_list:
            print(f"\n{'='*60}\n  {var_label}={level} | VUs={vus}\n{'='*60}")
            for env_name, env_path in ENVS.items():
                print(f"\n  --- [{env_name}] ---")
                md_path = os.path.join(env_path, "load_tests", uc_folder, "resultados",
                                       f"reporte_consolidado_{runs}runs_vus{vus}_nivel{level}.md")
                if os.path.exists(md_path):
                    print(f"    [SKIP] Reporte ya existe para {var_label}={level} VUs={vus}, reutilizando...")
                else:
                    restart_env_full(env_path)
                    run_script = os.path.join(env_path, "load_tests", "run_multiple_experiments.py")
                    print(f"    Ejecutando {runs} réplicas con {var_label}={level}, VUs={vus}...")
                    cmd = [sys.executable, run_script, "--test-script", script_rel,
                           "--vus", str(vus), "--runs", str(runs)]
                    env_vars = os.environ.copy()
                    env_vars[var_env_name] = str(level)
                    env_vars["LEVEL"] = str(level)
                    try:
                        p = subprocess.Popen(cmd, cwd=env_path,
                                             stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                             text=True, env=env_vars)
                        for line in p.stdout:
                            print(f"      {line.rstrip()}")
                        p.wait(timeout=1200)
                        if p.returncode != 0:
                            print(f"    [ERROR] Orquestador de K6 falló con código {p.returncode}.")
                            continue
                    except subprocess.TimeoutExpired:
                        print(f"    [TIMEOUT] Nivel {level} VUs {vus} tardó más de 20 min. Matando...")
                        p.kill(); p.wait()
                    except Exception as e:
                        print(f"    [ERROR] {e}")

                if os.path.exists(md_path):
                    res = parse_md_report(md_path)
                    if res:
                        full_data[env_name][(level, vus)] = res
                        avg = res['avg']
                        print(f"    [OK] Pet: {avg['peticiones']:.0f} | "
                              f"Lat: {avg['latencia']:.2f} ms | "
                              f"Fal: {avg['fallos']:.2f}% | "
                              f"Bloq: {avg.get('bloqueo_control', 0.0):.2f}%")

            combos_done.append((level, vus))
            excel_path = os.path.join(out_dir, f"{uc_id.replace('-','')}_Matriz.xlsx")
            try:
                export_matrix_excel(uc_id, var_label, full_data, levels, vus_list,
                                    excel_path, runs)
            except PermissionError:
                print(f"    [ERROR] Cerrá el Excel para poder sobrescribirlo.")
            except Exception as e:
                print(f"    [ERROR EXCEL] {e}")

    print("\n[+] Apagando entornos Docker...")
    for env_path in ENVS.values():
        try:
            subprocess.run(["docker-compose", "down"], cwd=env_path,
                          capture_output=True, timeout=60)
        except Exception:
            pass

    print("\n" + "=" * 70)
    print(f"  RESUMEN {uc_id} MATRIZ")
    print(f"  Celdas ejecutadas: {len(combos_done)} de {len(levels)*len(vus_list)}")
    print(f"  Archivos generados en: {out_dir}")
    print("=" * 70)
    return combos_done


def export_matrix_excel(uc_id, var_label, full_data, levels, vus_list, out_path, runs):
    """Exporta la matriz completa con una fila por combinación (level, vus, entorno, réplica)."""
    if not HAS_EXCEL:
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = uc_id.replace('-', '') + '_Matriz'
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color="000000")
    bdr = Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ttl_font = Font(bold=True, size=13, color="FFFFFF")
    hdr_font = Font(bold=True, color="FFFFFF", size=9)
    data_font = Font(color="000000", name="Calibri", size=9)
    fill_title = PatternFill("solid", fgColor="2F5597")
    fill_hdr = PatternFill("solid", fgColor="2F5597")
    fill_vuln = PatternFill("solid", fgColor="FFF2CC")
    fill_prot = PatternFill("solid", fgColor="D9E1F2")

    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A1:AF1")
    c = ws["A1"]
    c.value = f"{uc_id} MATRIZ - {var_label} × VUs - {now_str}"
    c.font = ttl_font; c.fill = fill_title; c.alignment = center

    HEADERS = [
        "Caso de Uso", "Entorno", "Tratamiento",
        var_label, "VUs", "Replica",
        "Fecha", "Hora",
        "Peticiones Totales", "Throughput (req/s)", "Latencia (ms)",
        "Fallos (%)", "Bloqueo Control (%)", "Disponibilidad (D)",
        "GW CPU Pico (%)", "GW CPU Prom (%)",
        "ms-usuarios CPU (%)", "ms-catalogo CPU (%)", "ms-resenas CPU (%)",
        "ms-ordenes CPU (%)", "Mongo CPU (%)", "Postgres CPU (%)",
        "GW RAM Pico (MiB)", "GW RAM Prom (MiB)",
        "ms-usuarios RAM (MiB)", "ms-catalogo RAM (MiB)", "ms-resenas RAM (MiB)",
        "ms-ordenes RAM (MiB)", "Mongo RAM (MiB)", "Postgres RAM (MiB)",
        "Tipo"
    ]
    for ci, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = hdr_font; cell.fill = fill_hdr; cell.alignment = center; cell.border = bdr

    row_idx = 4
    for level in levels:
        for vus in vus_list:
            for env_name in ["Vulnerable", "Protegido"]:
                is_vuln = (env_name == "Vulnerable")
                fill = fill_vuln if is_vuln else fill_prot
                tratam = "Linea Base Vulnerable" if is_vuln else "Hardening ISO 27001"
                fd = full_data.get(env_name, {}).get((level, vus), {})
                runs_list = fd.get("runs", [])
                if not runs_list:
                    cell = ws.cell(row=row_idx, column=1, value=uc_id)
                    cell.fill = fill; cell.alignment = center; cell.border = bdr
                    cell = ws.cell(row=row_idx, column=2, value=env_name)
                    cell.fill = fill; cell.alignment = center; cell.border = bdr
                    cell = ws.cell(row=row_idx, column=4, value=level)
                    cell.fill = fill; cell.alignment = center; cell.border = bdr
                    cell = ws.cell(row=row_idx, column=5, value=vus)
                    cell.fill = fill; cell.alignment = center; cell.border = bdr
                    cell = ws.cell(row=row_idx, column=len(HEADERS), value="Sin Datos")
                    cell.fill = fill; cell.alignment = center; cell.border = bdr
                    row_idx += 1
                    continue
                for run in runs_list:
                    tp = run.get('throughput_rps', run['peticiones'] / 60.0)
                    disp = run.get('disponibilidad',
                                   0 if (run['peticiones'] == 0 or run['fallos'] >= 50.0) else 1)
                    vals = [
                        uc_id, env_name, tratam, level, vus, f"R{run['idx']}",
                        run['fecha'], run['hora'],
                        round(run['peticiones'], 2), round(tp, 3),
                        round(run['latencia'], 2),
                        round(run['fallos'], 2),
                        round(run.get('bloqueo_control', 0.0), 2),
                        disp,
                        round(run.get('cpu_gw', 0), 2), round(run.get('cpu_gw_mean', 0), 2),
                        round(run.get('cpu_ms_usuarios', 0), 2),
                        round(run.get('cpu_ms_catalogo', 0), 2),
                        round(run.get('cpu_ms_resenas', 0), 2),
                        round(run.get('cpu_ms_ordenes', 0), 2),
                        round(run.get('cpu_mongo', 0), 2),
                        round(run.get('cpu_pg', 0), 2),
                        round(run.get('ram_gw', 0), 2), round(run.get('ram_gw_mean', 0), 2),
                        round(run.get('ram_ms_usuarios', 0), 2),
                        round(run.get('ram_ms_catalogo', 0), 2),
                        round(run.get('ram_ms_resenas', 0), 2),
                        round(run.get('ram_ms_ordenes', 0), 2),
                        round(run.get('ram_mongo', 0), 2),
                        round(run.get('ram_pg', 0), 2),
                        "Replica"
                    ]
                    for ci, v in enumerate(vals, 1):
                        cell = ws.cell(row=row_idx, column=ci, value=v)
                        cell.fill = fill; cell.alignment = center
                        cell.border = bdr; cell.font = data_font
                    row_idx += 1

    for ci in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 15
    ws.freeze_panes = "A4"
    wb.save(out_path)
    print(f"    [EXCEL] Matriz guardada: {out_path}")
