import os, re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PRUEBAS_DIR = os.path.join(BASE_DIR, "pruebas")
os.makedirs(PRUEBAS_DIR, exist_ok=True)

CONTAINERS = ['api-gateway', 'mongo-db', 'ms-catalogo', 'ms-ordenes', 'ms-resenas', 'ms-usuarios', 'postgres-db']
VUS_LIST = [1, 2, 3, 5, 8]

# Copy from panel_control.py
def parse_md_full_result(md_path):
    if not os.path.isfile(md_path): return {}
    with open(md_path, "r", encoding="utf-8") as f: txt = f.read()
    
    fecha = datetime.fromtimestamp(os.path.getmtime(md_path)).strftime("%d/%m/%Y")
    hora = datetime.fromtimestamp(os.path.getmtime(md_path)).strftime("%H:%M")
    
    run_metrics = {}
    lines = txt.split('\n')
    idx = 0
    while idx < len(lines):
        line = lines[idx]
        m = re.search(r'\|\s*\*\*Run (\d+)\*\*\s*\|\s*([\d.]+)\s*\|\s*([\d.]+)\s*ms\s*\|\s*([\d.]+)%\s*\|', line)
        if m:
            run_metrics[int(m.group(1))] = {
                "peticiones": float(m.group(2)),
                "latencia": float(m.group(3)),
                "fallos": float(m.group(4))
            }
        idx += 1
        
    cpu = {}
    for c in CONTAINERS:
        m = re.search(r'\|\s*\*\*' + re.escape(c) + r'\*\*\s*\|\s*([\d.]+)%\s*\|\s*([\d.]+)%\s*\|\s*([\d.]+)%\s*\|\s*\*\*([\d.]+)%\*\*', txt)
        if m:
            cpu[c] = {1: float(m.group(1)), 2: float(m.group(2)), 3: float(m.group(3))}
            
    ram = {}
    in_ram = False
    for line in lines:
        if "PICOS DE MEMORIA" in line: in_ram = True
        elif in_ram and "|" in line and "**" in line:
            for c in CONTAINERS:
                if f"**{c}**" in line:
                    parts = [p.strip() for p in line.split('|')]
                    mibs = []
                    for p in parts:
                        mm = re.search(r'\(([\d.]+)\s*MiB\)', p)
                        if mm: mibs.append(float(mm.group(1)))
                    if mibs:
                        ram[c] = {i+1: v for i,v in enumerate(mibs[:3])}
    
    avg = {"peticiones":0.0,"latencia":0.0,"fallos":0.0,"cpu_gw":0.0,"cpu_top":0.0}
    txt_clean = txt.replace('*', '')
    m_avg = re.search(r'PROMEDIO[^|]*\|\s*([\d.]+)\s*\|\s*([\d.]+)\s*ms\s*\|\s*([\d.]+)%', txt_clean)
    if m_avg:
        avg["peticiones"] = float(m_avg.group(1))
        avg["latencia"]   = float(m_avg.group(2))
        avg["fallos"]     = float(m_avg.group(3))
        
    gw_vals = list(cpu.get("api-gateway",{}).values())
    avg["cpu_gw"] = sum(gw_vals)/len(gw_vals) if gw_vals else 0.0
    
    max_cpu = 0.0
    for svc in CONTAINERS:
        if svc == "api-gateway": continue
        vals = list(cpu.get(svc,{}).values())
        if vals: max_cpu = max(max_cpu, max(vals))
    avg["cpu_top"] = max_cpu

    runs = []
    for n in sorted(run_metrics.keys()):
        rm = run_metrics[n]
        runs.append({
            "idx": n, "peticiones": rm["peticiones"], "latencia": rm["latencia"], "fallos": rm["fallos"],
            "fecha": fecha, "hora": hora,
            "cpu_gw": cpu.get("api-gateway",{}).get(n, 0.0),
            "cpu_ms_usuarios": cpu.get("ms-usuarios",{}).get(n, 0.0),
            "cpu_ms_catalogo": cpu.get("ms-catalogo",{}).get(n, 0.0),
            "cpu_ms_resenas": cpu.get("ms-resenas",{}).get(n, 0.0),
            "cpu_ms_ordenes": cpu.get("ms-ordenes",{}).get(n, 0.0),
            "cpu_mongo": cpu.get("mongo-db",{}).get(n, 0.0),
            "cpu_pg": cpu.get("postgres-db",{}).get(n, 0.0),
            "ram_gw": ram.get("api-gateway",{}).get(n, 0.0),
            "ram_ms_usuarios": ram.get("ms-usuarios",{}).get(n, 0.0),
            "ram_ms_catalogo": ram.get("ms-catalogo",{}).get(n, 0.0),
            "ram_ms_resenas": ram.get("ms-resenas",{}).get(n, 0.0),
            "ram_ms_ordenes": ram.get("ms-ordenes",{}).get(n, 0.0),
            "ram_mongo": ram.get("mongo-db",{}).get(n, 0.0),
            "ram_pg": ram.get("postgres-db",{}).get(n, 0.0),
        })

    return {"fecha": fecha, "hora": hora, "runs": runs, "avg": avg}

def export_to_excel(uc_name, level, full_data, avg_data, vus_list, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado 3 Runs"

    thin = openpyxl.styles.borders.Side(style="thin", color="000000")
    bdr = openpyxl.styles.borders.Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    ttl_font = Font(bold=True, size=14, color="FFFFFF")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    fill_title = PatternFill("solid", fgColor="2F5597")
    fill_hdr = PatternFill("solid", fgColor="2F5597")
    
    fill_vuln_a = PatternFill("solid", fgColor="FFF2CC")
    fill_vuln_b = PatternFill("solid", fgColor="FFE699")
    fill_prot_a = PatternFill("solid", fgColor="D9E1F2")
    fill_prot_b = PatternFill("solid", fgColor="B4C6E7")

    level_str = f" - Nivel {level}" if level else ""
    atk_name = f"{uc_name}{level_str}"
    
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws.merge_cells("A1:Y1")
    c = ws["A1"]
    c.value = f"{atk_name} - Exportado: {now_str}"
    c.font = ttl_font; c.fill = fill_title; c.alignment = center

    HEADERS = [
        "Caso de Uso", "Entorno", "Tratamiento", "Carga (VUs)", "Replica", "Fecha", "Hora",
        "Throughput (req/s)", "Latencia (ms)", "Fallos (%)", "GW CPU (%)", "ms-usuarios CPU (%)",
        "ms-catalogo CPU (%)", "ms-resenas CPU (%)", "ms-ordenes CPU (%)", "Mongo CPU (%)",
        "Postgres CPU (%)", "GW RAM (MiB)", "ms-usuarios RAM (MiB)", "ms-catalogo RAM (MiB)",
        "ms-resenas RAM (MiB)", "ms-ordenes RAM (MiB)", "Mongo RAM (MiB)", "Postgres RAM (MiB)", "Tipo"
    ]
    
    for ci, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = hdr_font; cell.fill = fill_hdr; cell.alignment = center; cell.border = bdr

    row_idx = 4
    for vus in vus_list:
        for env_name in ["Vulnerable", "Protegido"]:
            is_vuln = (env_name == "Vulnerable")
            fill_a = fill_vuln_a if is_vuln else fill_prot_a
            
            fd_env = full_data.get(env_name, {})
            fd_vus = fd_env.get(vus, {})
            runs = fd_vus.get("runs", [])
            
            for run in runs:
                vals = [
                    uc_name, env_name, "Linea Base Vulnerable" if is_vuln else "Hardening ISO 27001",
                    vus, f"R{run['idx']}", run['fecha'], run['hora'], run['peticiones'], run['latencia'],
                    run['fallos'], run['cpu_gw'], run['cpu_ms_usuarios'], run['cpu_ms_catalogo'], run['cpu_ms_resenas'],
                    run['cpu_ms_ordenes'], run['cpu_mongo'], run['cpu_pg'], run['ram_gw'], run['ram_ms_usuarios'],
                    run['ram_ms_catalogo'], run['ram_ms_resenas'], run['ram_ms_ordenes'], run['ram_mongo'], run['ram_pg'], "Replica"
                ]
                for ci, v in enumerate(vals, 1):
                    cell = ws.cell(row=row_idx, column=ci, value=v)
                    cell.fill = fill_a; cell.alignment = center; cell.border = bdr
                row_idx += 1

    wb.save(out_path)
    return True

def generate_all():
    from night_runner import CASES, VUS_LIST
    ENVS = [
        r"EntornoExperimental\Entorno_Vulnerable",
        r"EntornoExperimental\Entorno_Protegido"
    ]
    
    for case in CASES:
        uc_name = case["name"]
        uc_script = case["script"]
        uc_folder = os.path.dirname(uc_script)
        
        for level in case["levels"]:
            full_data = {"Vulnerable": {}, "Protegido": {}}
            avg_data = {"Vulnerable": {}, "Protegido": {}}
            
            for env in ENVS:
                env_key = "Vulnerable" if "controles" not in env else "Protegido"
                for vus in VUS_LIST:
                    level_str = f"_nivel{level}" if level is not None else ""
                    md_path = os.path.join(BASE_DIR, env, "load_tests", uc_folder, "resultados", f"reporte_consolidado_3runs_vus{vus}{level_str}.md")
                    
                    if os.path.exists(md_path):
                        res = parse_md_full_result(md_path)
                        full_data[env_key][vus] = res
                        avg_data[env_key][vus] = res.get("avg", {})
            
            # Excel
            level_str = f"_nivel{level}" if level is not None else ""
            out_excel = os.path.join(PRUEBAS_DIR, f"{uc_name.split(':')[0]}{level_str}_resultados.xlsx")
            export_to_excel(uc_name, level, full_data, avg_data, VUS_LIST, out_excel)
            
            # Simple Graphs (Throughput & Latencia)
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))
            fig.suptitle(f"{uc_name} {level_str.replace('_', ' ').title()}", fontsize=14, fontweight="bold")
            
            vuln_p = [avg_data["Vulnerable"].get(v, {}).get("peticiones", 0) for v in VUS_LIST]
            prot_p = [avg_data["Protegido"].get(v, {}).get("peticiones", 0) for v in VUS_LIST]
            vuln_l = [avg_data["Vulnerable"].get(v, {}).get("latencia", 0) for v in VUS_LIST]
            prot_l = [avg_data["Protegido"].get(v, {}).get("latencia", 0) for v in VUS_LIST]
            
            x = range(len(VUS_LIST))
            ax1.plot(x, vuln_p, marker='o', color='#E74C3C', label='Vulnerable')
            ax1.plot(x, prot_p, marker='s', color='#2ECC71', label='Protegido')
            ax1.set_xticks(x); ax1.set_xticklabels(VUS_LIST)
            ax1.set_title("Throughput (Peticiones/s)"); ax1.legend()
            
            ax2.plot(x, vuln_l, marker='o', color='#E74C3C', label='Vulnerable')
            ax2.plot(x, prot_l, marker='s', color='#2ECC71', label='Protegido')
            ax2.set_xticks(x); ax2.set_xticklabels(VUS_LIST)
            ax2.set_title("Latencia (ms)"); ax2.legend()
            
            plt.tight_layout()
            out_png = os.path.join(PRUEBAS_DIR, f"{uc_name.split(':')[0]}{level_str}_grafico.png")
            fig.savefig(out_png)
            plt.close(fig)
            print(f"[OK] Generado reporte y grafico para {uc_name}{level_str}")

if __name__ == "__main__":
    generate_all()
