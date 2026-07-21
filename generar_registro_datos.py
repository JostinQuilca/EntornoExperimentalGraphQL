# -*- coding: utf-8 -*-
"""
Construye la hoja Registro_Datos a partir de los reportes en disco.

Reproduce la parte de datos crudos del consolidado de la tesis: una fila por
replica, con las 7 variables dependientes y el consumo por microservicio.
No calcula estadistica ni genera graficas; para eso estan analisis_estadistico.py
y generar_todos_graficos.py.

Lee lo que haya en disco, asi que se puede ejecutar con la corrida a medias
para ver como va quedando.

Uso:
    python generar_registro_datos.py
    python generar_registro_datos.py --salida Registro_Datos.xlsx
"""
import argparse
import os
import sys
from datetime import datetime

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import uc_base
from experimento_completo import DISENO, REPLICAS, TRATAMIENTO, ruta_reporte

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# La exposicion estructural no sale de k6: se determino aparte comprobando si el
# esquema puede recorrerse por introspeccion. Es total en la linea base y nula
# con el control LNT-SEC-01 activo, sin valores intermedios.
EXPOSICION = {"Vulnerable": 100, "Protegido": 0}

GRUPOS = [
    ("IDENTIFICACIÓN", 8),
    ("MÉTRICAS DE RED Y SEGURIDAD", 6),
    ("CPU POR MICROSERVICIO (%)", 7),
    ("RAM POR MICROSERVICIO (MiB)", 7),
    ("", 1),
]

CABECERAS = [
    "Caso de Uso", "Control ISO", "Tipo de Amenaza", "Tratamiento",
    "Carga/Nivel", "Réplica", "Fecha", "Hora",
    "Latencia\n(ms)", "Throughput\n(req/s)", "Fallos\nEnet (%)",
    "Bloqueo\nControl (%)", "Disponi-\nbilidad (D)", "Exposición\nE (%)",
    "CPU\nGW (%)", "CPU\nUsuarios (%)", "CPU\nCatálogo (%)", "CPU\nReseñas (%)",
    "CPU\nÓrdenes (%)", "CPU\nMongo (%)", "CPU\nPostgres (%)",
    "RAM\nGW (MiB)", "RAM\nUsuarios (MiB)", "RAM\nCatálogo (MiB)", "RAM\nReseñas (MiB)",
    "RAM\nÓrdenes (MiB)", "RAM\nMongo (MiB)", "RAM\nPostgres (MiB)",
    "Observaciones",
]

# Campos de cada replica tal como los devuelve uc_base.parse_md_report.
CAMPOS_METRICA = [
    "latencia", "throughput_rps", "fallos", "bloqueo_control", "disponibilidad",
]
CAMPOS_CPU = [
    "cpu_gw", "cpu_ms_usuarios", "cpu_ms_catalogo", "cpu_ms_resenas",
    "cpu_ms_ordenes", "cpu_mongo", "cpu_pg",
]
CAMPOS_RAM = [
    "ram_gw", "ram_ms_usuarios", "ram_ms_catalogo", "ram_ms_resenas",
    "ram_ms_ordenes", "ram_mongo", "ram_pg",
]


def recolectar(replicas):
    """Recorre el diseno y devuelve una fila por replica encontrada en disco."""
    filas, faltantes = [], []
    for uc_id, uc in DISENO.items():
        for env_name in ("Vulnerable", "Protegido"):
            for vus, nivel in uc["escenarios"]:
                md = ruta_reporte(env_name, uc, vus, nivel, replicas)
                datos = uc_base.parse_md_report(md) if os.path.isfile(md) else None
                if not datos or not datos.get("runs"):
                    faltantes.append((uc_id, env_name, vus, nivel))
                    continue

                for r in datos["runs"]:
                    obs = ""
                    # Una latencia de 0 ms no es una medicion rapida: es que el
                    # gateway no contesto. Se marca para que no pase inadvertida.
                    if r["latencia"] == 0:
                        obs = "Sin respuesta del entorno; réplica no válida"

                    fila = [
                        uc["nombre"], uc["iso"], uc["amenaza"], TRATAMIENTO[env_name],
                        uc["carga"](vus, nivel), f"R{r['idx']}", r["fecha"], r["hora"],
                    ]
                    fila += [r[c] for c in CAMPOS_METRICA[:4]]
                    fila += [r["disponibilidad"], EXPOSICION[env_name]]
                    fila += [r[c] for c in CAMPOS_CPU]
                    fila += [r[c] for c in CAMPOS_RAM]
                    fila.append(obs)
                    filas.append(fila)
    return filas, faltantes


def escribir(filas, salida):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro_Datos"

    fuente = "Arial"
    azul = "1F4E78"
    borde = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Fila 1: titulo
    ws.cell(1, 1, "REGISTRO DE DATOS CRUDOS — 5 réplicas por escenario — "
                  "7 variables — consumo por microservicio")
    ws.cell(1, 1).font = Font(name=fuente, size=12, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=azul)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CABECERAS))
    ws.row_dimensions[1].height = 24

    # Fila 2: grupos de columnas
    col = 1
    for nombre, ancho in GRUPOS:
        if nombre:
            c = ws.cell(2, col, nombre)
            c.font = Font(name=fuente, size=9, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2E75B6")
            c.alignment = centro
            c.border = borde
            if ancho > 1:
                ws.merge_cells(start_row=2, start_column=col,
                               end_row=2, end_column=col + ancho - 1)
        col += ancho
    ws.row_dimensions[2].height = 18

    # Fila 3: cabeceras
    for i, h in enumerate(CABECERAS, 1):
        c = ws.cell(3, i, h)
        c.font = Font(name=fuente, size=8, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=azul)
        c.alignment = centro
        c.border = borde
    ws.row_dimensions[3].height = 38

    # Datos
    relleno = {
        "Línea Base Vulnerable": PatternFill("solid", fgColor="FCE4E4"),
        "Hardening ISO 27001":   PatternFill("solid", fgColor="E2EFDA"),
    }
    for n, fila in enumerate(filas, start=4):
        for i, v in enumerate(fila, 1):
            c = ws.cell(n, i, v)
            c.font = Font(name=fuente, size=8)
            c.alignment = centro if i != len(fila) else Alignment(
                horizontal="left", vertical="center", wrap_text=True)
            c.border = borde
            c.fill = relleno.get(fila[3], PatternFill())
            if i in (9, 10) or 15 <= i <= 28:
                c.number_format = "0.00"

    anchos = {1: 30, 2: 10, 3: 18, 4: 22, 5: 12, 6: 8, 7: 11, 8: 8, 29: 34}
    for i in range(1, len(CABECERAS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(i, 10)

    ws.freeze_panes = "A4"

    # Nota al pie: de donde sale cada valor que no viene medido de k6.
    nota = len(filas) + 5
    ws.cell(nota, 1, "Notas:")
    ws.cell(nota, 1).font = Font(name=fuente, size=9, bold=True)
    for i, txt in enumerate([
        "Exposición E%: 100 en la línea base y 0 con el control activo. No la mide k6; se determinó "
        "comprobando aparte si el esquema puede recorrerse por introspección.",
        "Disponibilidad D: 1 si el gateway procesó la carga; 0 si no hubo peticiones o los fallos "
        "llegaron al 50%.",
        "Throughput: peticiones completadas sobre la ventana de 60 s de cada réplica.",
        "CPU y RAM: pico por contenedor durante la réplica, tomado de docker stats.",
    ], start=1):
        c = ws.cell(nota + i, 1, f"• {txt}")
        c.font = Font(name=fuente, size=8)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.merge_cells(start_row=nota + i, start_column=1,
                       end_row=nota + i, end_column=len(CABECERAS))

    wb.save(salida)


def main():
    ap = argparse.ArgumentParser(description="Genera la hoja Registro_Datos desde los reportes en disco")
    ap.add_argument("--salida", default=None, help="ruta del .xlsx de salida")
    ap.add_argument("--replicas", type=int, default=REPLICAS)
    args = ap.parse_args()

    salida = args.salida or os.path.join(
        BASE_DIR, f"Registro_Datos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    filas, faltantes = recolectar(args.replicas)
    if not filas:
        sys.exit("[ERROR] No se encontro ningun reporte en disco. Ejecuta antes experimento_completo.py")

    escribir(filas, salida)

    esperadas = sum(len(uc["escenarios"]) for uc in DISENO.values()) * 2 * args.replicas
    invalidas = sum(1 for f in filas if f[-1])
    print(f"[OK] {salida}")
    print(f"  Filas escritas   : {len(filas)} de {esperadas} esperadas")
    if faltantes:
        print(f"  Escenarios sin reporte ({len(faltantes)}):")
        for uc_id, env_name, vus, nivel in faltantes[:15]:
            print(f"    - {uc_id} | {env_name} | VUs={vus} Nivel={nivel}")
        if len(faltantes) > 15:
            print(f"    ... y {len(faltantes) - 15} mas")
    if invalidas:
        print(f"  [AVISO] {invalidas} replicas sin respuesta del entorno, marcadas en Observaciones.")


if __name__ == "__main__":
    main()
